import docx
import traceback

try:
    doc = docx.Document(r'c:\Users\00003\OneDrive\Desktop\areef\دراسة_جدوى_عريف.docx')
    
    with open(r'c:\Users\00003\OneDrive\Desktop\areef\extracted.txt', 'w', encoding='utf-8') as f:
        for p in doc.paragraphs:
            f.write(p.text + '\n')
        
        for i, table in enumerate(doc.tables):
            f.write(f'\n--- TABLE {i+1} ---\n')
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                f.write('\t'.join(cells) + '\n')
    
    print('done')
except Exception as e:
    traceback.print_exc()
